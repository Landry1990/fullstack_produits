"""
Tests pour api/views/purge.py (PurgeViewSet — maintenance superadmin).

Couvre : tables, preview, export ZIP, purge (avec mot de passe),
backup/restore (validation), produits_count, import_produits (validation),
import_status, export_produits, purge_produits, download_rapport,
changelog, update_status, run_update.
"""
import io
import zipfile
from datetime import timedelta
from decimal import Decimal
from unittest.mock import patch

from django.core.cache import cache
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient

from ..models import (
    AuditLog,
    Caisse,
    Facture,
    FactureProduit,
    MouvementCaisse,
    Produit,
)
from .factories import TestDataFactory as F

ADMIN_PWD = 'adminpass123'


class PurgeBase(TestCase):
    def setUp(self):
        self.client_api = APIClient()
        self.admin = F.create_superuser(password=ADMIN_PWD)
        self.client_api.force_authenticate(user=self.admin)
        cache.clear()


# ── Accès ─────────────────────────────────────────────────────────────────────

class PurgeAccessTest(PurgeBase):
    def test_non_staff_refuse(self):
        user = F.create_user()
        client = APIClient()
        client.force_authenticate(user=user)
        resp = client.get(reverse('maintenance-tables'))
        self.assertEqual(resp.status_code, 403)

    def test_anonyme_refuse(self):
        resp = APIClient().get(reverse('maintenance-tables'))
        self.assertIn(resp.status_code, (401, 403))


# ── Tables / Preview / Export / Purge ─────────────────────────────────────────

class PurgeTablesTest(PurgeBase):
    def test_tables_liste_categories(self):
        resp = self.client_api.get(reverse('maintenance-tables'))
        self.assertEqual(resp.status_code, 200)
        keys = {t['key'] for t in resp.data}
        for expected in ('factures', 'commandes', 'caisse', 'audit_logs'):
            self.assertIn(expected, keys)

    def test_preview_sans_tables_400(self):
        resp = self.client_api.post(reverse('maintenance-preview'), {})
        self.assertEqual(resp.status_code, 400)

    def test_preview_compte_parents_et_enfants(self):
        facture = F.create_facture(status='VAL', total_ttc=Decimal('1000'))
        produit = F.create_produit()
        F.create_facture_produit(facture, produit, quantity=1)
        F.create_caisse(facture, 1000)

        resp = self.client_api.post(reverse('maintenance-preview'), {
            'tables': ['factures'],
        }, format='json')
        self.assertEqual(resp.status_code, 200)
        row = resp.data[0]
        self.assertEqual(row['key'], 'factures')
        self.assertEqual(row['count'], 1)
        child_labels = {c['label']: c['count'] for c in row['children']}
        self.assertEqual(child_labels['Lignes facture'], 1)
        self.assertEqual(child_labels['Paiements caisse'], 1)

    def test_preview_filtre_par_dates(self):
        vieille = F.create_facture(status='VAL')
        Facture.objects.filter(id=vieille.id).update(
            date=timezone.now() - timedelta(days=400)
        )
        F.create_facture(status='VAL')  # récente

        resp = self.client_api.post(reverse('maintenance-preview'), {
            'tables': ['factures'],
            'date_to': (timezone.now() - timedelta(days=300)).date().isoformat(),
        }, format='json')
        self.assertEqual(resp.data[0]['count'], 1)

    def test_preview_table_inconnue_ignoree(self):
        resp = self.client_api.post(reverse('maintenance-preview'), {
            'tables': ['table_qui_existe_pas'],
        }, format='json')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data, [])

    def test_export_zip_csv(self):
        facture = F.create_facture(status='VAL')
        F.create_caisse(facture, 500)

        resp = self.client_api.post(reverse('maintenance-export'), {
            'tables': ['factures'],
        }, format='json')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/zip')
        zf = zipfile.ZipFile(io.BytesIO(resp.content))
        noms = zf.namelist()
        self.assertIn('factures.csv', noms)
        # Enfants exportés aussi
        self.assertTrue(any('paiements' in n for n in noms))

    def test_export_sans_tables_400(self):
        resp = self.client_api.post(reverse('maintenance-export'), {})
        self.assertEqual(resp.status_code, 400)


class PurgeExecuteTest(PurgeBase):
    def test_sans_password_400(self):
        resp = self.client_api.post(reverse('maintenance-purge'), {
            'tables': ['factures'],
        }, format='json')
        self.assertEqual(resp.status_code, 400)

    def test_mauvais_password_403(self):
        F.create_facture(status='VAL')
        resp = self.client_api.post(reverse('maintenance-purge'), {
            'tables': ['factures'],
            'password': 'mauvais_mdp',
        }, format='json')
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(Facture.objects.count(), 1)  # rien supprimé

    def test_non_superuser_403_meme_avec_bon_password(self):
        staff = F.create_user(password=ADMIN_PWD)
        staff.is_staff = True
        staff.save()
        client = APIClient()
        client.force_authenticate(user=staff)
        F.create_facture(status='VAL')
        resp = client.post(reverse('maintenance-purge'), {
            'tables': ['factures'],
            'password': ADMIN_PWD,
        }, format='json')
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(Facture.objects.count(), 1)

    def test_purge_supprime_et_logue(self):
        facture = F.create_facture(status='VAL')
        F.create_facture_produit(facture, F.create_produit(), quantity=1)
        F.create_caisse(facture, 800)

        resp = self.client_api.post(reverse('maintenance-purge'), {
            'tables': ['factures'],
            'password': ADMIN_PWD,
        }, format='json')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['results'][0]['deleted'], 1)
        # CASCADE : lignes + paiements supprimés
        self.assertEqual(Facture.objects.count(), 0)
        self.assertEqual(FactureProduit.objects.count(), 0)
        self.assertEqual(Caisse.objects.count(), 0)
        # Trace d'audit
        self.assertTrue(
            AuditLog.objects.filter(model_name='Purge').exists()
        )

    def test_purge_respecte_la_plage_de_dates(self):
        vieille = F.create_facture(status='VAL')
        Facture.objects.filter(id=vieille.id).update(
            date=timezone.now() - timedelta(days=400)
        )
        recente = F.create_facture(status='VAL')

        resp = self.client_api.post(reverse('maintenance-purge'), {
            'tables': ['factures'],
            'date_to': (timezone.now() - timedelta(days=300)).date().isoformat(),
            'password': ADMIN_PWD,
        }, format='json')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['results'][0]['deleted'], 1)
        self.assertEqual(Facture.objects.filter(id=recente.id).count(), 1)


# ── Produits : count / import / export / purge ────────────────────────────────

class PurgeProduitsTest(PurgeBase):
    def test_produits_count(self):
        F.create_produit()
        F.create_produit()
        resp = self.client_api.get(reverse('maintenance-produits-count'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['count'], Produit.objects.count())

    def test_import_sans_fichier_400(self):
        resp = self.client_api.post(reverse('maintenance-import-produits'), {})
        self.assertEqual(resp.status_code, 400)

    def test_import_mauvaise_extension_400(self):
        fichier = io.BytesIO(b'col\nval')
        fichier.name = 'produits.txt'
        resp = self.client_api.post(
            reverse('maintenance-import-produits'), {'file': fichier},
        )
        self.assertEqual(resp.status_code, 400)

    def test_import_deja_en_cours_409(self):
        cache.set('import_produits_running', 'job123', timeout=60)
        fichier = io.BytesIO(b'nom,prix_vente\nTest,100')
        fichier.name = 'produits.csv'
        resp = self.client_api.post(
            reverse('maintenance-import-produits'), {'file': fichier},
        )
        self.assertEqual(resp.status_code, 409)

    def test_import_status_idle(self):
        resp = self.client_api.get(reverse('maintenance-import-status'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['status'], 'idle')

    def test_import_status_job_inconnu_404(self):
        resp = self.client_api.get(
            reverse('maintenance-import-status'), {'job_id': 'inconnu'},
        )
        self.assertEqual(resp.status_code, 404)

    def test_export_produits_xlsx(self):
        F.create_produit(name='Produit Export', selling_price=250)
        resp = self.client_api.get(reverse('maintenance-export-produits'))
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=4).value, 'nom')
        noms = [ws.cell(row=r, column=4).value for r in range(2, ws.max_row + 1)]
        self.assertIn('Produit Export', noms)

    def test_purge_produits_sans_password_400(self):
        resp = self.client_api.post(reverse('maintenance-purge-produits'), {})
        self.assertEqual(resp.status_code, 400)

    def test_purge_produits_mauvais_password_403(self):
        F.create_produit()
        resp = self.client_api.post(reverse('maintenance-purge-produits'), {
            'password': 'faux',
        }, format='json')
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(Produit.objects.count(), 1)

    def test_purge_produits_tout_supprime(self):
        F.create_produit()
        F.create_produit()
        resp = self.client_api.post(reverse('maintenance-purge-produits'), {
            'password': ADMIN_PWD,
        }, format='json')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(Produit.objects.count(), 0)

    def test_purge_produits_sans_ventes_conserve_lies(self):
        produit_vendu = F.create_produit(name='Vendu')
        F.create_produit(name='Jamais vendu')
        facture = F.create_facture(status='VAL')
        F.create_facture_produit(facture, produit_vendu, quantity=1)

        resp = self.client_api.post(reverse('maintenance-purge-produits'), {
            'password': ADMIN_PWD,
            'sans_ventes': True,
        }, format='json')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['conserves'], 1)
        self.assertTrue(Produit.objects.filter(name='Vendu').exists())
        self.assertFalse(Produit.objects.filter(name='Jamais vendu').exists())


# ── Rapports d'import / changelog / update ────────────────────────────────────

class MaintenanceDiversTest(PurgeBase):
    def test_download_rapport_nom_invalide_400(self):
        for mauvais in ('../settings.py', 'a/b.txt', ''):
            resp = self.client_api.get(
                reverse('maintenance-download-rapport'), {'file': mauvais},
            )
            self.assertEqual(resp.status_code, 400)

    def test_download_rapport_inexistant_404(self):
        resp = self.client_api.get(
            reverse('maintenance-download-rapport'), {'file': 'inexistant.txt'},
        )
        self.assertEqual(resp.status_code, 404)

    def test_changelog_introuvable_404(self):
        import tempfile
        with tempfile.TemporaryDirectory() as tmpdir:
            with override_settings(BASE_DIR=tmpdir):
                resp = self.client_api.get(reverse('maintenance-changelog'))
        self.assertEqual(resp.status_code, 404)

    def test_changelog_lu(self):
        import tempfile
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmpdir:
            Path(tmpdir, 'CHANGELOG.md').write_text(
                '# Changelog — Fullstack Produits\n\n## Section récente\ncontenu',
                encoding='utf-8',
            )
            with override_settings(BASE_DIR=tmpdir):
                resp = self.client_api.get(reverse('maintenance-changelog'))
        self.assertEqual(resp.status_code, 200)
        self.assertIn('Section récente', resp.data['latest'])

    def test_update_status_idle(self):
        resp = self.client_api.get(reverse('maintenance-update-status'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['status'], 'idle')

    def test_run_update_mauvais_password_403(self):
        resp = self.client_api.post(reverse('maintenance-run-update'), {
            'password': 'faux',
        })
        self.assertEqual(resp.status_code, 403)

    def test_run_update_script_absent_404(self):
        import tempfile
        with tempfile.TemporaryDirectory() as tmpdir:
            with override_settings(BASE_DIR=tmpdir):
                resp = self.client_api.post(reverse('maintenance-run-update'), {
                    'password': ADMIN_PWD,
                })
        self.assertEqual(resp.status_code, 404)


# ── Backup / Restore via commandes (validation) ───────────────────────────────

class MaintenanceBackupRestoreTest(PurgeBase):
    def test_restore_sans_fichier_400(self):
        resp = self.client_api.post(reverse('maintenance-restore'), {
            'password': ADMIN_PWD,
        })
        self.assertEqual(resp.status_code, 400)

    def test_restore_sans_password_400(self):
        fichier = io.BytesIO(b'gz data')
        fichier.name = 'backup.sql.gz'
        resp = self.client_api.post(reverse('maintenance-restore'), {
            'file': fichier,
        })
        self.assertEqual(resp.status_code, 400)

    def test_restore_mauvais_password_403(self):
        fichier = io.BytesIO(b'gz data')
        fichier.name = 'backup.sql.gz'
        resp = self.client_api.post(reverse('maintenance-restore'), {
            'file': fichier,
            'password': 'faux',
        })
        self.assertEqual(resp.status_code, 403)

    def test_backup_erreur_commande_500(self):
        with patch('api.views.purge.call_command', side_effect=Exception('pg absent')):
            resp = self.client_api.post(reverse('maintenance-backup'))
        self.assertEqual(resp.status_code, 500)
        self.assertIn('Erreur', resp.data['detail'])
