"""
Tests pour api/views/rapports/finance.py (RapportFinanceMixin, ~7% couvert avant).

Couvre : rapport mensuel/par dates (JSON+PDF), CA multi-annuel, TVA vendus,
export comptable CSV, remises (JSON+Excel), livre de caisse Excel,
détail marges, stats marges, rapport dynamique, export Sage i7,
rapport fiscal mensuel.
"""
import io
import json
from datetime import timedelta
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from ..models import (
    Caisse,
    Commande,
    FactureProduitAllocation,
    PharmacySettings,
)
from .factories import TestDataFactory as F


def _dates():
    today = timezone.localtime(timezone.now()).date()
    return today.isoformat(), today.isoformat()


class RapportsFinanceBase(TestCase):
    """Données communes : 1 facture validée + ligne + paiement + mouvement caisse."""

    def setUp(self):
        self.client_api = APIClient()
        self.admin = F.create_superuser()
        self.client_api.force_authenticate(user=self.admin)

        self.produit = F.create_produit(
            name='Produit Finance', stock=100, cost_price=200, selling_price=500,
        )
        self.produit.tva = Decimal('19.25')
        self.produit.pmp = Decimal('200')
        self.produit.save()

        self.client_obj = F.create_client(name='Client Finance')
        self.today = timezone.localtime(timezone.now()).date()
        self.db, self.df = self.today.isoformat(), self.today.isoformat()

        # Facture validée : 2 x 500 = 1000 brut, remise 100, TTC 900
        # NB : les signaux recalculent les totaux à l'ajout de la ligne
        # (TTC 900 → HT 755 + TVA 145 au taux 19.25%)
        self.facture = F.create_facture(
            client=self.client_obj,
            status='VAL',
            numero_facture='FAC-FIN-001',
            total_ht=Decimal('1000.00'),
            total_tva=Decimal('0.00'),
            total_ttc=Decimal('900.00'),
            remise=Decimal('100.00'),
            validated_by=self.admin,
            created_by=self.admin,
        )
        self.ligne = F.create_facture_produit(
            self.facture, self.produit, quantity=2,
            selling_price=Decimal('500'), tva=Decimal('19.25'),
        )
        self.facture.refresh_from_db()
        self.paiement = F.create_caisse(
            self.facture, 900, mode_paiement='especes', user=self.admin,
        )
        F.create_mouvement_caisse(self.admin, 'ENTREE', 5000, 'Fond de caisse')
        F.create_mouvement_caisse(self.admin, 'SORTIE', 500, 'Course')


# ── Permissions ───────────────────────────────────────────────────────────────

class PermissionsFinanceTest(TestCase):
    def test_non_authentifie_refuse(self):
        resp = APIClient().get(reverse('rapports-rapport-mensuel'), {'mois': '2026-01'})
        self.assertIn(resp.status_code, (401, 403))

    def test_utilisateur_sans_menu_rapports_refuse(self):
        user = F.create_user()
        client = APIClient()
        client.force_authenticate(user=user)
        resp = client.get(reverse('rapports-rapport-mensuel'), {'mois': '2026-01'})
        self.assertEqual(resp.status_code, 403)


# ── Rapports JSON de base ─────────────────────────────────────────────────────

class RapportsJsonTest(RapportsFinanceBase):
    def test_rapport_mensuel_sans_mois_400(self):
        resp = self.client_api.get(reverse('rapports-rapport-mensuel'))
        self.assertEqual(resp.status_code, 400)

    def test_rapport_mensuel_format_invalide_400(self):
        resp = self.client_api.get(reverse('rapports-rapport-mensuel'), {'mois': 'pasunmois'})
        self.assertEqual(resp.status_code, 400)

    def test_rapport_mensuel_ok(self):
        mois = self.today.strftime('%Y-%m')
        resp = self.client_api.get(reverse('rapports-rapport-mensuel'), {'mois': mois})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('ca', resp.data)

    def test_rapport_par_dates_sans_dates_400(self):
        resp = self.client_api.get(reverse('rapports-rapport-par-dates'))
        self.assertEqual(resp.status_code, 400)

    def test_rapport_par_dates_ok(self):
        resp = self.client_api.get(
            reverse('rapports-rapport-par-dates'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['date_debut'], self.db)
        self.assertEqual(resp.data['date_fin'], self.df)


# ── CA multi-annuel ───────────────────────────────────────────────────────────

class RapportCAMultiAnnuelTest(RapportsFinanceBase):
    def test_structure_13_lignes(self):
        resp = self.client_api.get(reverse('rapports-rapport-ca-multi-annuel'))
        self.assertEqual(resp.status_code, 200)
        data = resp.data
        # 12 mois + ligne total
        self.assertEqual(len(data), 13)
        self.assertEqual(data[0]['Mois'], 'january')
        self.assertEqual(data[-1]['Mois'], 'total_general')
        annee = self.today.year
        self.assertIn(f'{annee}_total', data[0])
        # CA du mois courant = total_ttc de la facture (900)
        mois_courant = data[self.today.month - 1]
        self.assertEqual(float(mois_courant[f'{annee}_total']), 900.0)

    def test_sans_facture_retourne_liste_vide(self):
        FactureProduitAllocation.objects.all().delete()
        from ..models import Facture
        Facture.objects.all().update(status='ANN')
        resp = self.client_api.get(reverse('rapports-rapport-ca-multi-annuel'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data, [])


# ── TVA vendus ────────────────────────────────────────────────────────────────

class RapportTvaVendusTest(RapportsFinanceBase):
    def test_sans_dates_400(self):
        resp = self.client_api.get(reverse('rapports-rapport-tva-vendus'))
        self.assertEqual(resp.status_code, 400)

    def test_dates_invalides_400(self):
        resp = self.client_api.get(
            reverse('rapports-rapport-tva-vendus'),
            {'date_debut': 'xx', 'date_fin': 'yy'},
        )
        self.assertEqual(resp.status_code, 400)

    def test_ligne_tva_calculee(self):
        resp = self.client_api.get(
            reverse('rapports-rapport-tva-vendus'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        row = next(r for r in resp.data if r['produit'] == 'Produit Finance')
        self.assertEqual(row['quantite'], 2)
        self.assertEqual(row['taux_tva'], '19.25 %')
        # TTC ligne = 2 * 500 = 1000 ; TVA = 1000 * 19.25 / 119.25 ≈ 161
        self.assertEqual(row['total_ttc'], 1000)
        self.assertAlmostEqual(float(row['montant_tva']), 161.4, delta=1.0)


# ── Export comptable CSV ──────────────────────────────────────────────────────

class ExportComptableCsvTest(RapportsFinanceBase):
    def test_csv_contenu(self):
        resp = self.client_api.get(
            reverse('rapports-export-comptable-csv'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'text/csv')
        content = resp.content.decode('utf-8-sig')
        self.assertIn('Espèces', content)          # mode de paiement affiché
        self.assertIn('900,00', content)           # total TTC format français
        self.assertIn(self.facture.numero_facture or str(self.facture.id), content)

    def test_periode_trop_longue_400(self):
        debut = (self.today - timedelta(days=800)).isoformat()
        resp = self.client_api.get(
            reverse('rapports-export-comptable-csv'),
            {'date_debut': debut, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 400)


# ── Remises ───────────────────────────────────────────────────────────────────

class RapportRemisesTest(RapportsFinanceBase):
    def test_rapport_remises_par_vendeur(self):
        resp = self.client_api.get(
            reverse('rapports-rapport-remises'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        row = next(r for r in resp.data if r['user_id'] == self.admin.id)
        self.assertEqual(float(row['remise_globale']), 100.0)
        self.assertEqual(row['nb_factures'], 1)
        self.assertGreater(float(row['total_remise']), 0)

    def test_rapport_remises_details(self):
        resp = self.client_api.get(
            reverse('rapports-rapport-remises-details'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        row = resp.data[0]
        self.assertEqual(row['numero_facture'], self.facture.numero_facture)
        self.assertEqual(row['remise_globale'], 100.0)

    def test_facture_sans_remise_exclue_des_details(self):
        F.create_facture(client=self.client_obj, status='VAL', total_ttc=Decimal('500'))
        resp = self.client_api.get(
            reverse('rapports-rapport-remises-details'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        # Une seule facture a une remise > 0
        self.assertEqual(len(resp.data), 1)

    def test_remises_excel(self):
        resp = self.client_api.get(
            reverse('rapports-rapport-remises-excel'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws.title, 'Remises')

    def test_remises_details_excel(self):
        resp = self.client_api.get(
            reverse('rapports-rapport-remises-details-excel'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.active.title, 'Détails Remises')


# ── Livre de caisse Excel ─────────────────────────────────────────────────────

class LivreCaisseExcelTest(RapportsFinanceBase):
    def test_sans_dates_400(self):
        resp = self.client_api.get(reverse('rapports-livre-caisse-excel'))
        self.assertEqual(resp.status_code, 400)

    def test_excel_deux_feuilles_et_totaux(self):
        resp = self.client_api.get(
            reverse('rapports-livre-caisse-excel'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(len(wb.sheetnames), 2)

        ws = wb['Livre de Caisse']
        rows = list(ws.iter_rows(values_only=True))
        # Ligne TOTAL GÉNÉRAL : espèces=900, entrées=5000, sorties=500
        total_row = next(r for r in rows if r[0] == 'TOTAL GÉNÉRAL')
        self.assertEqual(total_row[1], 900.0)          # Espèces
        self.assertEqual(total_row[11], 5000.0)        # Entrées manuelles
        self.assertEqual(total_row[12], 500.0)         # Sorties manuelles
        self.assertEqual(total_row[13], 5400.0)        # Solde = 900+5000-500


# ── Détail marges ─────────────────────────────────────────────────────────────

class RapportDetailMargesTest(RapportsFinanceBase):
    def setUp(self):
        super().setUp()
        self.lot = F.create_stock_lot(self.produit, quantity=50, price_cost=200)
        FactureProduitAllocation.objects.create(
            facture_produit=self.ligne, stock_lot=self.lot,
            quantity=2, cost_price=Decimal('200'), selling_price=Decimal('500'),
        )

    def _url(self, **extra):
        params = {'date_debut': self.db, 'date_fin': self.df}
        params.update(extra)
        return self.client_api.get(reverse('rapports-rapport-detail-marges'), params)

    def test_ligne_allouee_marge(self):
        resp = self._url()
        self.assertEqual(resp.status_code, 200)
        rows = resp.data.get('results', resp.data) if isinstance(resp.data, dict) else resp.data
        row = next(r for r in rows if r['produit'] == 'Produit Finance')
        self.assertEqual(row['lot'], 'LOT-TEST-001')
        # Vente nette = (500-0)*2 * ratio(900/1000) = 900 ; achat = 200*2 = 400
        self.assertEqual(row['mt_achat'], 400.0)
        self.assertAlmostEqual(row['mt_vente'], 900.0, delta=1.0)
        self.assertGreater(row['marge'], 0)

    def test_ligne_sans_lot_fallback_pmp(self):
        produit2 = F.create_produit(name='Sans Lot', cost_price=100, selling_price=300)
        produit2.pmp = Decimal('100')
        produit2.save()
        F.create_facture_produit(self.facture, produit2, quantity=1, selling_price=Decimal('300'))
        resp = self._url()
        rows = resp.data.get('results', resp.data) if isinstance(resp.data, dict) else resp.data
        row = next(r for r in rows if r['produit'] == 'Sans Lot')
        self.assertEqual(row['lot'], 'SANS LOT')
        self.assertEqual(row['cout_achat'], 100.0)

    def test_grouper_par_produit(self):
        resp = self._url(grouper_par='produit')
        rows = resp.data.get('results', resp.data) if isinstance(resp.data, dict) else resp.data
        row = next(r for r in rows if r['produit'] == 'Produit Finance')
        self.assertIn('statut', row)
        self.assertEqual(row['statut'], 'OK')
        self.assertIn('nb_ventes', row)

    def test_filtre_marge_negative(self):
        # Produit vendu à perte via allocation à coût élevé
        prod_perte = F.create_produit(name='En Perte', cost_price=600, selling_price=100)
        ligne2 = F.create_facture_produit(self.facture, prod_perte, quantity=1, selling_price=Decimal('100'))
        lot2 = F.create_stock_lot(prod_perte, quantity=10, price_cost=600, lot_name='LOT-CHER')
        FactureProduitAllocation.objects.create(
            facture_produit=ligne2, stock_lot=lot2,
            quantity=1, cost_price=Decimal('600'), selling_price=Decimal('100'),
        )
        resp = self._url(filtre_marge='negative')
        rows = resp.data.get('results', resp.data) if isinstance(resp.data, dict) else resp.data
        self.assertTrue(all(r['taux_marge'] < 0 for r in rows))
        self.assertTrue(any(r['produit'] == 'En Perte' for r in rows))


# ── Stats marges ──────────────────────────────────────────────────────────────

class StatsMargesTest(RapportsFinanceBase):
    def test_ligne_jour_et_total(self):
        resp = self.client_api.get(
            reverse('rapports-stats-marges'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(resp.data), 2)          # 1 jour + TOTAL
        jour = resp.data[0]
        self.assertEqual(jour['ca'], 900.0)
        self.assertEqual(jour['nb_ventes'], 1)
        self.assertEqual(resp.data[-1]['date'], 'TOTAL')

    def test_periode_vide_retourne_liste_vide(self):
        debut = (self.today + timedelta(days=30)).isoformat()
        fin = (self.today + timedelta(days=31)).isoformat()
        resp = self.client_api.get(
            reverse('rapports-stats-marges'),
            {'date_debut': debut, 'date_fin': fin},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data, [])


# ── Rapport dynamique ─────────────────────────────────────────────────────────

class RapportDynamiqueTest(RapportsFinanceBase):
    def _url(self, **extra):
        params = {'date_debut': self.db, 'date_fin': self.df}
        params.update(extra)
        return self.client_api.get(reverse('rapports-rapport-dynamique'), params)

    def test_dates_requises_sauf_produits(self):
        resp = self._url(source='ventes')
        resp = self.client_api.get(reverse('rapports-rapport-dynamique'),
                                   {'source': 'ventes'})
        self.assertEqual(resp.status_code, 400)

    def test_source_ventes(self):
        resp = self._url(
            source='ventes',
            fields='date,facture,produit,quantite,prix_vente,total_ht,marge',
        )
        self.assertEqual(resp.status_code, 200)
        row = resp.data[0]
        self.assertEqual(row['Produit'], 'Produit Finance')
        self.assertEqual(row['Quantité'], 2)
        self.assertIn('Marge Brute', row)

    def test_source_achats(self):
        commande = F.create_commande(status='CLOT')
        commande.date_cloture = timezone.now()
        commande.save()
        F.create_commande_produit(commande, self.produit, quantity=10,
                                  price_cost=Decimal('200'))
        resp = self._url(
            source='achats',
            fields='date,fournisseur,produit,quantite,cout_achat,total_ht',
        )
        self.assertEqual(resp.status_code, 200)
        row = next(r for r in resp.data if r.get('Produit') == 'Produit Finance')
        self.assertEqual(row['Total HT'], 2000.0)

    def test_source_stock(self):
        F.create_stock_lot(self.produit, quantity=50, price_cost=200)
        resp = self._url(
            source='stock',
            fields='produit,lot,quantite,cout_achat,total_ht',
        )
        self.assertEqual(resp.status_code, 200)
        row = resp.data[0]
        self.assertEqual(row['Lot'], 'LOT-TEST-001')
        self.assertEqual(row['Quantité'], 50)

    def test_group_by_agrege(self):
        resp = self._url(
            source='ventes',
            fields='produit,quantite,total_ht',
            group_by='Produit',
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn('_count', resp.data[0])

    def test_sort_by_colonne(self):
        resp = self._url(
            source='ventes', fields='produit,quantite',
            sort_by='Quantité', sort_order='asc',
        )
        self.assertEqual(resp.status_code, 200)
        quantites = [r['Quantité'] for r in resp.data]
        self.assertEqual(quantites, sorted(quantites))


# ── Export Sage i7 ────────────────────────────────────────────────────────────

class ExportSageTest(RapportsFinanceBase):
    def test_csv_ecritures_vente_et_reglement(self):
        resp = self.client_api.get(
            reverse('rapports-export-sage-i7'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        content = resp.content.decode('utf-8-sig')
        # Journal VT : débit client 411100 + crédit ventes 701100
        self.assertIn('VT', content)
        self.assertIn('411100', content)
        self.assertIn('701100', content)
        # Paiement espèces → journal CA compte 571100
        self.assertIn('CA', content)
        self.assertIn('571100', content)
        # Libellé règlement avec le mode affiché (régression get_mode_paiement_display)
        self.assertIn('Regl Espèces', content)
        self.assertNotIn('Regl None', content)

    def test_dates_invalides_400(self):
        resp = self.client_api.get(
            reverse('rapports-export-sage-i7'),
            {'date_debut': 'xx', 'date_fin': 'yy'},
        )
        self.assertEqual(resp.status_code, 400)


# ── Rapport fiscal mensuel ────────────────────────────────────────────────────

class RapportFiscalTest(RapportsFinanceBase):
    def setUp(self):
        super().setUp()
        self.ps = PharmacySettings.objects.create(
            regime_fiscal='REEL',
            mode_imposition='DROIT_COMMUN',
            taux_accompte_reel=Decimal('2.00'),
            taux_cac=Decimal('10.00'),
            taux_precompte_reel=Decimal('1.00'),
        )

    def test_structure_et_calcul_accompte(self):
        resp = self.client_api.get(
            reverse('rapports-rapport-fiscal-mensuel'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.data
        self.assertEqual(data['regime_fiscal'], 'REEL')
        self.assertEqual(data['chiffre_affaires']['ca_ttc'], 900)
        # Droit commun réel : accompte = CA HT (755) * 2% ; CAC = 10% de l'accompte
        self.assertEqual(data['accompte']['base_imposition'], 755)
        self.assertEqual(data['accompte']['accompte_base'], 15)
        self.assertEqual(data['accompte']['accompte_cac'], 2)
        self.assertEqual(data['accompte']['accompte_total'], 17)

    def test_mode_marge_administree(self):
        self.ps.mode_imposition = 'MARGE_ADMINISTREE'
        self.ps.taux_marge_brute = Decimal('14.00')
        self.ps.save()
        resp = self.client_api.get(
            reverse('rapports-rapport-fiscal-mensuel'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.data
        # Marge brute = CA HT (755) - coût ventes (200*2=400) = 355 ; accompte = 355*14% ≈ 50
        self.assertEqual(data['accompte']['base_label'], 'Marge brute')
        self.assertEqual(data['accompte']['base_imposition'], 355)
        self.assertEqual(data['accompte']['accompte_base'], 50)

    def test_sans_settings_400(self):
        self.ps.delete()
        resp = self.client_api.get(
            reverse('rapports-rapport-fiscal-mensuel'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 400)

    def test_sans_dates_400(self):
        resp = self.client_api.get(reverse('rapports-rapport-fiscal-mensuel'))
        self.assertEqual(resp.status_code, 400)


# ── PDF / Excel général ───────────────────────────────────────────────────────

class ExportsPdfTest(RapportsFinanceBase):
    def test_rapport_mensuel_pdf(self):
        mois = self.today.strftime('%Y-%m')
        resp = self.client_api.get(
            reverse('rapports-rapport-mensuel-pdf'), {'mois': mois},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')

    def test_rapport_par_dates_pdf(self):
        resp = self.client_api.get(
            reverse('rapports-rapport-par-dates-pdf'),
            {'date_debut': self.db, 'date_fin': self.df},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')

    def test_rapport_general_excel(self):
        mois = self.today.strftime('%Y-%m')
        resp = self.client_api.get(
            reverse('rapports-rapport-general-excel'), {'mois': mois},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])

    def test_excel_general_coherent_avec_ventes_partagees(self):
        from ..services.ventes_lots import collect_sales_records
        from ..utils.dates import day_start
        from ..utils_doclang import T
        from ..views.rapports.excel_general import _collect_data, build_rapport_general_excel, _fmt, _pct
        from ..views.rapports.excel_general_extra import collect_extra_data

        fournisseur = F.create_fournisseur(is_divers=True)
        divers = F.create_produit(name='Divers sans lot', cost_price=90, pmp=40, selling_price=100, fournisseur=fournisseur)
        F.create_facture_produit(self.facture, divers, quantity=2, selling_price=Decimal('100'), discount=Decimal('10'), tva=0)
        lot = F.create_stock_lot(self.produit, price_cost=Decimal('150'))
        FactureProduitAllocation.objects.create(
            facture_produit=self.ligne, stock_lot=lot, quantity=2,
            selling_price=Decimal('777'), cost_price=Decimal('150'),
        )
        F.create_caisse(self.facture, 1, user=self.admin)
        for status, active, payment_status in [('PAY', False, 'completee'), ('VAL', True, 'en_attente')]:
            excluded = F.create_facture(status=status, is_active=active)
            F.create_facture_produit(excluded, self.produit, quantity=50)
            payment = F.create_caisse(excluded, 1, user=self.admin)
            Caisse.objects.filter(pk=payment.pk).update(statut=payment_status)
        start = day_start(self.today.replace(day=1))
        end = day_start((self.today.replace(day=28) + timedelta(days=4)).replace(day=1))
        factor = Decimal('1080') / Decimal('1180')
        expected_ht = Decimal('1000') * factor / Decimal('1.1925') + Decimal('180') * factor
        expected_margin = expected_ht - Decimal('380')
        records = collect_sales_records(invoices=type(self.facture).objects.filter(pk=self.facture.pk))
        self.assertAlmostEqual(float(sum(r['vente_ht'] for r in records)), float(expected_ht))
        settings, _ = PharmacySettings.objects.get_or_create(pk=1)
        settings.mode_imposition = 'MARGE_ADMINISTREE'
        settings.taux_marge_brute = Decimal('10')
        settings.taux_cac = Decimal('10')
        settings.save()
        for lang in ('fr', 'en'):
            data = _collect_data(start, end, lang=lang)
            self.assertEqual(data['ca']['nb_ventes'], 1)
            self.assertAlmostEqual(float(data['ca']['ca_ttc']), 1080)
            self.assertAlmostEqual(float(data['ca']['ca_ht']), float(expected_ht))
            self.assertAlmostEqual(float(data['marge']['marge_brute']), float(expected_margin))
            self.assertAlmostEqual(sum(r['marge_totale'] for r in data['daily_rows']), float(expected_margin))
            self.assertAlmostEqual(float(sum(r['marge'] for r in data['top_rows'])), float(expected_margin))
            self.assertEqual(sum(r['quantite'] for r in data['top_rows']), 4)
            self.assertAlmostEqual(float(data['divers_ca']['ca_ttc']), float(Decimal('180') * factor))
            self.assertAlmostEqual(float(data['marge']['marge_divers']), float(Decimal('180') * factor - 80))
            self.assertEqual(data['accompte_fiscal']['base_imposition'], round(expected_margin))
            extra = collect_extra_data(start, end, lang=lang)
            self.assertAlmostEqual(float(extra['evolution']['marge_cur']), float(expected_margin))
            response = build_rapport_general_excel(start, end, 'Test', lang=lang)
            wb = load_workbook(io.BytesIO(response.content), read_only=True)
            summary = {row[0]: row[1] for row in wb[T(lang, 'xs_sheet_synthese')].iter_rows(values_only=True) if len(row) > 1}
            self.assertEqual(summary[T(lang, 'xs_r_marge_totale')], _fmt(expected_margin))
            self.assertEqual(summary[T(lang, 'xs_r_taux_total')], _pct(expected_margin / expected_ht * 100))
            daily = list(wb[T(lang, 'xs_sheet_ca_marges')].iter_rows(values_only=True))
            self.assertEqual(daily[8][7], _fmt(expected_margin))
            self.assertEqual(daily[-1][8], _pct(expected_margin / expected_ht * 100))
            top = list(wb[T(lang, 'xs_sheet_top')].iter_rows(values_only=True))
            for cells, product in zip(top[8:], data['top_rows']):
                self.assertEqual(cells[8], _fmt(product['marge']))
                self.assertEqual(cells[9], _pct(product['taux_marge']))
            fiscal = {row[0]: row[1] for row in wb[T(lang, 'xs_sheet_fiscal')].iter_rows(values_only=True) if len(row) > 1}
            self.assertEqual(fiscal[T(lang, 'fis_base_marge')], _fmt(data['accompte_fiscal']['base_imposition']))
            wb.close()

    def test_excel_general_bornes_mensuelles_et_periode_vide(self):
        from ..utils.dates import day_start
        from ..views.rapports.excel_general import _collect_data
        from ..views.rapports.excel_general_extra import collect_extra_data

        start = day_start(self.today.replace(day=1))
        end = day_start((self.today.replace(day=28) + timedelta(days=4)).replace(day=1))
        type(self.facture).objects.filter(pk=self.facture.pk).update(date=start - timedelta(microseconds=1))
        future = F.create_facture(status='PAY')
        F.create_facture_produit(future, self.produit, quantity=100)
        type(future).objects.filter(pk=future.pk).update(date=end)
        data = _collect_data(start, end, lang='fr')
        self.assertEqual(data['ca']['ca_ttc'], 0)
        self.assertEqual(data['ca']['nb_ventes'], 0)
        self.assertEqual(data['marge']['marge_brute'], 0)
        self.assertEqual(data['daily_rows'], [])
        self.assertEqual(data['top_rows'], [])
        extra = collect_extra_data(start, end, lang='fr')
        self.assertEqual(extra['evolution']['ca_cur'], 0)
        self.assertEqual(extra['evolution']['marge_cur'], 0)
        self.assertEqual(extra['evolution']['ca_prev'], 900)
        self.assertAlmostEqual(float(extra['evolution']['marge_prev']), float(Decimal('900') / Decimal('1.1925') - 400))

    def test_excel_general_base_fiscale_ca_et_marge(self):
        from ..utils.dates import day_start
        from ..views.rapports.excel_general import _collect_data

        start = day_start(self.today.replace(day=1))
        end = day_start((self.today.replace(day=28) + timedelta(days=4)).replace(day=1))
        settings, _ = PharmacySettings.objects.get_or_create(pk=1)
        settings.mode_imposition = 'MARGE_ADMINISTREE'
        settings.taux_marge_brute = Decimal('10')
        settings.taux_cac = Decimal('10')
        settings.save()
        data = _collect_data(start, end, lang='fr')
        self.assertEqual(data['accompte_fiscal']['base_imposition'], 355)
        self.assertEqual(data['accompte_fiscal']['accompte_base'], 35)
        self.assertEqual(data['accompte_fiscal']['accompte_total'], 39)
        settings.mode_imposition = 'DROIT_COMMUN'
        settings.regime_fiscal = 'REEL'
        settings.taux_accompte_reel = Decimal('2')
        settings.save()
        data = _collect_data(start, end, lang='fr')
        self.assertEqual(data['accompte_fiscal']['base_imposition'], 755)
        self.assertEqual(data['accompte_fiscal']['accompte_base'], 15)
        self.assertEqual(data['accompte_fiscal']['accompte_total'], 17)

    def test_rapport_general_excel_mois_invalide_400(self):
        resp = self.client_api.get(
            reverse('rapports-rapport-general-excel'), {'mois': 'bad'},
        )
        self.assertEqual(resp.status_code, 400)
