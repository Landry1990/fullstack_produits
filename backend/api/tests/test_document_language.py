from datetime import datetime
from io import BytesIO
from unittest.mock import patch

import openpyxl
from django.test import TestCase

from api.models import InvoiceSettings, PharmacySettings
from api.services.invoice_pdf import generate_invoice_pdf
from api.utils_doclang import DOC_STRINGS, T, format_doc_date, get_document_language
from api.views.rapports.excel_general import build_rapport_general_excel

from .factories import TestDataFactory


class DocumentLanguageHelperTests(TestCase):
    def test_pharmacy_locale_controls_document_language(self):
        settings = PharmacySettings.objects.create(locale='en-US')
        self.assertEqual(get_document_language(), 'en')

        settings.locale = 'fr-FR'
        settings.save(update_fields=['locale'])
        self.assertEqual(get_document_language(), 'fr')

    def test_missing_settings_falls_back_to_french(self):
        self.assertEqual(get_document_language(), 'fr')

    def test_translation_catalogs_have_identical_keys(self):
        self.assertSetEqual(set(DOC_STRINGS['fr']), set(DOC_STRINGS['en']))

    def test_translation_and_unknown_key_fallbacks(self):
        self.assertEqual(T('en', 'col_produit'), 'Product')
        self.assertEqual(T('fr', 'col_produit'), 'Produit')
        self.assertEqual(T('en', 'unknown_document_key'), 'unknown_document_key')
        self.assertEqual(T('de', 'col_produit'), 'Produit')

    def test_document_dates_follow_selected_language(self):
        value = datetime(2026, 9, 22, 14, 35)
        self.assertEqual(format_doc_date(value, 'fr'), '22/09/2026')
        self.assertEqual(format_doc_date(value, 'en'), '09/22/2026')
        self.assertEqual(format_doc_date(value, 'en', with_time=True), '09/22/2026 14:35')


class DocumentGenerationSmokeTests(TestCase):
    def setUp(self):
        PharmacySettings.objects.create(
            pharmacy_name='Test Pharmacy',
            address='Test Address',
            city='Douala',
            phone='600000000',
            locale='en-US',
        )

    def test_general_excel_uses_document_language(self):
        response = build_rapport_general_excel(
            datetime(2026, 9, 1),
            datetime(2026, 9, 30),
            'September 2026',
            lang='en',
        )

        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(BytesIO(response.content), read_only=True)
        self.assertIn('Summary', workbook.sheetnames)
        self.assertIn('Sales & Margins', workbook.sheetnames)
        self.assertNotIn('Synthèse', workbook.sheetnames)

        summary = workbook['Summary']
        visible_values = {
            value
            for row in summary.iter_rows(min_row=1, max_row=15, values_only=True)
            for value in row
            if isinstance(value, str)
        }
        self.assertIn('Indicator', visible_values)
        self.assertTrue(any('MONTHLY SUMMARY' in value for value in visible_values))

    @patch('api.services.invoice_pdf.valider_licence_systeme', return_value=(False, None, None))
    def test_invoice_pdf_generates_in_english(self, _mock_licence):
        facture = TestDataFactory.create_facture(total_ttc=1000)
        invoice_settings = InvoiceSettings.objects.create(
            company_name='Test Pharmacy',
            company_address='Test Address',
        )

        pdf = generate_invoice_pdf(facture, invoice_settings, lang='en').getvalue()

        self.assertGreater(len(pdf), 500)
        self.assertTrue(pdf.startswith(b'%PDF'))
